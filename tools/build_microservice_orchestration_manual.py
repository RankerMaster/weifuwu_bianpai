from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs"
DOCX_PATH = OUTPUT_DIR / "微服务容器编排使用说明.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(85, 95, 110)
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_cm):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), "9360")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths_cm):
                cell.width = Cm(widths_cm[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("微服务容器编排系统使用说明")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    r = subtitle.add_run("适用于 WPF 微服务编排管理程序、Docker 镜像构建、远程容器运维与策略预部署。")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(11)
    r.font.color.rgb = MUTED

    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    set_table_widths(table, [4.0, 12.51])
    rows = [
        ("项目位置", str(ROOT)),
        ("生成日期", datetime.now().strftime("%Y-%m-%d")),
        ("主要程序", "WpfApp1 / 微服务编排管理"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        set_cell_shading(row.cells[0], HEADER_FILL)
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(value)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_key_value_table(doc, title, rows):
    doc.add_heading(title, level=3)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_widths(table, [4.75, 11.76])
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "说明"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, HEADER_FILL)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    set_table_widths(table, [4.75, 11.76])


def add_command_block(doc, commands):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_widths(table, [16.51])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("\n".join(commands))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def build_doc():
    OUTPUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_heading("1. 系统概述", level=1)
    doc.add_paragraph(
        "本系统是一个 Windows WPF 桌面端微服务容器编排工具，面向已部署 Docker 的边缘终端或远程服务器。"
        "它通过本机 Docker CLI、SSH Docker Context 和远程 SSH 命令读取设备运行状态，完成镜像构建、镜像同步、容器创建、容器启停、模型文件写入、算法包部署、压力测试和策略预部署。"
    )
    add_bullets(doc, [
        "镜像和容器感知：读取设备、镜像、容器、CPU、内存、磁盘和运行数量。",
        "微服务镜像编排：基于本地基础镜像和程序包生成自定义镜像，并同步到目标设备。",
        "微服务容器编排：在目标设备上创建、启动、停止、删除容器，查看日志，执行模型和算法部署。",
        "容器策略优化：根据镜像大小、使用次数、使用间隔和内存预算生成预部署镜像方案，并一键推送。",
    ])

    doc.add_heading("2. 使用前准备", level=1)
    add_key_value_table(doc, "2.1 运行环境", [
        ("操作系统", "Windows，建议安装 .NET Framework 4.8 运行环境。"),
        ("桌面程序", "运行 WpfApp1.exe，开发环境可打开 WpfApp1.sln 编译运行。"),
        ("Docker CLI", "本机需要安装 Docker Desktop 或 Docker CLI，且命令行可执行 docker。"),
        ("远程设备", "目标服务器需安装 Docker，并允许 SSH 连接。"),
        ("网络", "本机到目标设备的 SSH 端口需可达，常见为 22 端口。"),
        ("权限", "SSH 用户需能执行 docker 命令；若 Docker 需 sudo，需确保远程命令可正常执行。"),
    ])
    doc.add_heading("2.2 推荐目录约定", level=3)
    add_bullets(doc, [
        "程序包目录：dockerbuild_workspace\\program_package，用于放置业务引擎、依赖库、Xenomai、压力测试等 tar.gz 包。",
        "Dockerfile 模板目录：dockerbuild_workspace\\template_dockfile，用于保存 ai、bh、db、img、stress 等构建模板。",
        "临时构建目录：dockerbuild_workspace\\tempBuild_image，系统创建镜像时会写入 Dockerfile、程序包和导出的镜像 tar。",
        "策略日志目录：dockerbuild_workspace\\strategy_logs，用于保存不同设备的策略 CSV 数据。",
        "测试输出目录：outputs，用于保存压力测试、通信延迟和本说明文档等结果文件。",
    ])
    doc.add_heading("2.3 Docker Context 检查", level=3)
    doc.add_paragraph("系统会优先识别 SSH 类型的 Docker Context。首次使用前建议在 PowerShell 中确认 Docker 可用：")
    add_command_block(doc, [
        "docker version",
        "docker context ls",
        "docker context inspect <context-name>",
        "docker --context <context-name> images",
        "docker --context <context-name> ps -a",
    ])
    doc.add_paragraph("如果没有预建 Context，也可以在程序中通过“读取设备信息”输入 SSH 用户名、目标 IP 和密码，程序会在当前会话中使用 SSH 密码通道读取目标 Docker 数据。")

    doc.add_heading("3. 首次连接设备", level=1)
    add_numbers(doc, [
        "启动 WpfApp1.exe，进入“微服务编排管理”主界面。",
        "在左侧选择“镜像和容器感知”。",
        "点击“读取设备信息”。",
        "在弹窗中输入 SSH 用户名、通信目标 IP 地址和 SSH 密码。",
        "等待程序完成密码校验、Docker 运行信息拉取和界面刷新。",
        "连接成功后，页面会显示已连接设备数量、镜像数量、已启动容器数量、CPU/内存/磁盘占用和镜像列表。",
    ])
    doc.add_paragraph("连接成功后，设备会被命名为“设备1、设备2 ...”，并映射到对应的 SSH Docker Context 或当前会话 SSH 密码通道。后续镜像编排、容器编排和策略优化都依赖该设备映射。")

    doc.add_heading("4. 镜像和容器感知", level=1)
    add_bullets(doc, [
        "资源仪表：展示当前设备或全部设备的 CPU、内存、磁盘占用。",
        "设备选择：可在下拉框中切换“全部设备”或单台设备。",
        "镜像列表：展示 Name、Tag、Created、Size、Image ID 等字段。",
        "右键镜像：支持修改 Name、修改 Tag、删除镜像等操作。",
        "刷新：重新读取远程 Docker 镜像、容器和资源指标。",
    ])
    doc.add_paragraph("饼图区域支持鼠标悬浮或点击扇区查看具体资源占用说明。若页面为空，通常表示尚未读取设备信息，或目标设备 Docker/SSH 不可达。")

    doc.add_heading("5. 微服务镜像编排", level=1)
    doc.add_paragraph("该模块用于把本地基础镜像、业务程序包和 Dockerfile 组合成可部署镜像，并同步到远程设备。")
    add_key_value_table(doc, "5.1 界面区域", [
        ("基础镜像列表", "来自本机 Docker 的普通镜像，不包含已生成的 custom_image 镜像。"),
        ("镜像列表", "主要显示本地已生成或可部署的自定义镜像。"),
        ("程序包列表", "读取 dockerbuild_workspace\\program_package 下的业务包。"),
        ("设备选择", "选择全部设备或单台设备，用于查看与部署。"),
    ])
    doc.add_heading("5.2 新建镜像", level=3)
    add_numbers(doc, [
        "选择“微服务镜像编排”。",
        "在“基础镜像列表”中选中一个基础镜像。",
        "在“程序包列表”中选择一个或多个程序包。",
        "点击“新建镜像”。",
        "在 Dockerfile 编辑窗口确认或调整自动生成的 Dockerfile。",
        "点击确认后等待 docker build、docker image save 和界面刷新完成。",
    ])
    doc.add_paragraph("构建成功后，系统会在 tempBuild_image 目录生成 Dockerfile、镜像 tar 文件和 meta 文件；镜像标签通常形如 local:custom_image_yyyyMMddHHmmss。")
    doc.add_heading("5.3 部署镜像到设备", level=3)
    add_numbers(doc, [
        "先完成“读取设备信息”，确保目标设备可用。",
        "在“镜像列表”中选中一个或多个镜像。",
        "点击“部署”。",
        "在弹出的设备选择窗口中选择目标设备。",
        "等待系统完成镜像同步。同步过程会优先复用本机镜像；目标设备不存在时会执行导出、传输和导入。"
    ])
    doc.add_heading("5.4 下载、发布、删除", level=3)
    add_bullets(doc, [
        "下载镜像：从本机下载目录选择镜像 tar 后执行 docker load，并尝试识别镜像标签。",
        "发布：对选中的镜像执行发布流程，适合把镜像推送到目标或指定位置。",
        "删除镜像：可从基础镜像列表或镜像列表中选择镜像后删除；删除前确认镜像未被关键容器使用。",
    ])
    doc.add_heading("5.5 算法部署", level=3)
    doc.add_paragraph("算法部署用于选择本地 JAR 包，并更新 dockerbuild_workspace\\tempBuild_image 下的 Dockerfile COPY 行。适合在构建镜像前替换算法包。")

    doc.add_heading("6. 微服务容器编排", level=1)
    doc.add_paragraph("该模块直接面向目标设备上的容器运行管理。进入页面后，先选择设备，再在容器列表中进行操作。")
    add_key_value_table(doc, "6.1 容器列表字段", [
        ("名称", "中文显示名，便于业务人员识别。"),
        ("镜像", "容器使用的镜像名称与标签。"),
        ("状态", "running、exited、paused 等 Docker 状态。"),
        ("CPU 核数量", "容器 CPU 限制或宿主机核数显示。"),
        ("CPU/Memory", "来自 docker stats 的实时占用。"),
        ("Port(s)", "端口映射。"),
        ("Size", "容器大小或限制信息。"),
        ("详情/ID", "Docker 状态详情和容器 ID。"),
    ])
    doc.add_heading("6.2 创建容器", level=3)
    add_numbers(doc, [
        "点击“创建容器”。",
        "在弹窗中选择镜像并填写容器名称。",
        "设置 CPU 核数，最小 0.5。",
        "设置内存限制，单位 MB，必须为正整数并且是 4 MB 的倍数。",
        "填写端口映射；至少填写一个 Host 端口，且端口范围为 1-65535。",
        "确认后系统执行 Docker 创建/启动流程，并写入 CASS 相关配置。",
    ])
    doc.add_paragraph("容器名称只能包含字母、数字、下划线、点和横杠。同一设备上不能创建同名容器。")
    doc.add_heading("6.3 批量操作", level=3)
    add_bullets(doc, [
        "启动容器：选中一个或多个容器后点击“启动容器”。启动超时时间较长，适合等待业务服务初始化。",
        "停止容器：选中容器后点击“停止容器”。",
        "删除容器：选中容器后点击“删除容器”。删除前请确认业务数据已保存。",
        "查看日志：选中容器后点击“查看日志”，用于排查服务启动失败或运行异常。",
        "右键修改容器名称：在容器行上右键，选择“修改容器名称”。",
    ])
    doc.add_heading("6.4 模型部署", level=3)
    doc.add_paragraph("模型部署用于把本地 JSON 文件写入选中容器的固定路径 /home/CASS-Simulator/test.json。远程设备场景下，程序会先通过 SCP 上传到宿主机临时目录，再执行 docker cp 写入容器。")
    add_numbers(doc, [
        "在容器列表中选中目标容器，可多选。",
        "点击“模型部署”。",
        "选择本地 JSON 文件。",
        "等待系统完成容器校验、远端上传、docker cp 写入和结果确认。",
    ])
    doc.add_heading("6.5 压力测试", level=3)
    doc.add_paragraph("压力测试用于评估目标设备可同时创建或运行的容器数量。系统会按指定镜像、名称前缀、最大尝试数量、CPU、内存、端口起点等参数批量创建容器。")
    add_bullets(doc, [
        "模式 create：只创建容器，不启动。",
        "模式 run：创建并启动容器。",
        "Cleanup：测试后清理压力测试容器。",
        "BatchSize：按批次记录日志和进度。",
        "结果：系统会显示成功创建数量，并写入容器操作日志。",
    ])

    doc.add_heading("7. 容器策略优化", level=1)
    doc.add_paragraph("策略优化模块用于根据本地 Docker 镜像数据和内存预算生成预部署镜像列表。系统会调用 cuckoo_Train1.py，结合镜像大小、使用次数和使用间隔等数据选择预部署镜像。")
    add_numbers(doc, [
        "选择“容器策略优化”。",
        "选择目标设备。",
        "点击“读取磁盘上限”，确认目标设备 Docker 根目录或宿主机磁盘容量。",
        "输入“边缘终端最大内存阈值”，单位 MB。",
        "调整百分比滑块，得到策略预算。",
        "点击“预部署镜像生成”。",
        "检查生成的预部署镜像列表。",
        "点击“一键部署”，在确认窗口中选择需要部署的镜像并执行推送。",
    ])
    doc.add_paragraph("若 Python 策略脚本执行失败，请检查 cuckoo_Train1.py 路径、Python 环境、image_plan_data.csv 和镜像列表是否可读。")

    doc.add_heading("8. 测试与记录", level=1)
    add_key_value_table(doc, "8.1 推荐测试项", [
        ("连通性测试", "读取设备信息，确认 SSH 密码校验和 Docker 信息读取成功。"),
        ("镜像构建测试", "选择基础镜像和程序包，新建 custom_image 并确认 tar 输出。"),
        ("镜像同步测试", "把一个小镜像部署到远程设备并刷新远端镜像列表。"),
        ("容器生命周期测试", "创建、启动、查看日志、停止、删除同一个测试容器。"),
        ("压力测试", "从较小 Max 值开始，例如 10 或 50，确认无端口冲突后逐步加大。"),
        ("通信延迟测试", "可参考项目中的“容器间最小通信延迟测试方法.md”，使用 docker network 和 ping 统计 RTT_min/2。"),
    ])
    doc.add_heading("8.2 容器间最小通信延迟参考命令", level=3)
    add_command_block(doc, [
        "docker pull alpine:latest",
        "docker network create latency-test",
        "docker run -dit --name latency-c1 --network latency-test alpine sh",
        "docker run -dit --name latency-c2 --network latency-test alpine sh",
        "docker exec latency-c1 sh -c \"apk add --no-cache iputils\"",
        "docker exec latency-c1 ping -c 500 latency-c2",
    ])
    doc.add_paragraph("读取 rtt min/avg/max/mdev 中的 min 值，最小单向通信时延可近似按 RTT_min / 2 计算。")

    doc.add_heading("9. 常见问题处理", level=1)
    add_key_value_table(doc, "9.1 故障排查表", [
        ("读取设备失败", "检查 SSH 用户名、密码、IP、网络、防火墙和目标机 SSH 服务。"),
        ("Docker 命令无输出", "确认本机 docker 可执行，目标设备 Docker 服务运行，SSH 用户具备 Docker 权限。"),
        ("页面显示暂无设备", "先点击“读取设备信息”；若仍为空，执行 docker context ls 和 docker --context <name> ps -a 排查。"),
        ("镜像构建失败", "检查基础镜像是否存在、程序包是否可读、Dockerfile COPY 路径是否正确。"),
        ("部署超时", "检查镜像体积、网络速度、目标磁盘空间；可手工执行 docker image inspect、docker load 或 docker images 验证。"),
        ("容器创建失败", "检查容器名重复、端口占用、内存限制是否为 4 MB 倍数、镜像是否存在。"),
        ("模型部署失败", "确认选中容器存在，JSON 文件可读，目标路径 /home/CASS-Simulator/test.json 所在目录在容器内可写。"),
        ("策略生成失败", "确认本机 Docker 镜像可读取，Python 策略脚本路径有效，策略预算大于 0。"),
    ])

    doc.add_heading("10. 安全与维护建议", level=1)
    add_bullets(doc, [
        "生产设备上执行删除镜像、删除容器、压力测试前，应先确认没有承载关键业务。",
        "压力测试建议使用专用镜像和独立名称前缀，测试后开启清理或手工删除残留容器。",
        "SSH 密码仅在当前程序会话中用于远程命令和部署通道；退出程序后需重新读取设备信息。",
        "定期清理 tempBuild_image 中的临时构建文件和过期镜像 tar，避免占满本机磁盘。",
        "远程设备磁盘空间不足时，先执行 docker system df 分析，再按需清理 dangling 镜像、停止容器或旧镜像。",
        "修改 Dockerfile 模板和程序包后，建议先在测试设备验证镜像构建和容器启动，再部署到正式设备。",
    ])

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("微服务容器编排系统使用说明")
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build_doc())
